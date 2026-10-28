"""Create the Day 6-7 Word deliverables for the APL Logistics project."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path("deliverables")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(13 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    # Word's built-in Title style can carry a blue lower border. Remove it so
    # document titles remain plain black typography with whitespace only.
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)


def add_title_block(document: Document, title: str, subtitle: str) -> None:
    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_paragraph.add_run(title)
    subtitle_paragraph = document.add_paragraph()
    subtitle_paragraph.paragraph_format.space_after = Pt(15)
    run = subtitle_paragraph.add_run(subtitle)
    run.bold = True
    run.font.size = Pt(12)
    subtitle_paragraph.paragraph_format.space_after = Pt(18)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_paragraph(text, style=f"Heading {level}")


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header_cells = table.rows[0].cells
    for index, label in enumerate(headers):
        cell = header_cells[index]
        cell.text = label
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
            run.font.size = Pt(10)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.text = value
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, "EDF3F8")
            if widths:
                cell.width = Inches(widths[index])
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(10)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def write_research_paper() -> Path:
    document = Document()
    configure_document(document)
    add_title_block(
        document,
        "Delivery Performance Delay Risk and Logistics Efficiency Analysis",
        "Research paper on global supply chain operations",
    )

    add_heading(document, "Abstract")
    document.add_paragraph(
        "This study examines delivery performance and delay risk in 180,519 historical global supply chain shipment records. "
        "The analysis compares actual shipping duration with scheduled duration, classifies deliveries as early, on time, or delayed, and evaluates performance by shipping mode, region, market, country, and customer segment. "
        "The dataset shows that 57.28 percent of shipments were delayed according to the calculated delivery gap, while 42.72 percent were on time or early. "
        "The main operational concern is the First Class shipping mode, which was delayed for every recorded shipment. "
        "The resulting Streamlit dashboard gives stakeholders a reusable view of delay exposure and delivery performance rather than a one-time spreadsheet analysis."
    )
    keywords = document.add_paragraph()
    keywords.add_run("Keywords: ").bold = True
    keywords.add_run("supply chain analytics, delivery performance, delay risk, logistics dashboard, Streamlit")

    add_heading(document, "1 Introduction")
    document.add_paragraph(
        "Large logistics operations manage shipments across markets, regions, product categories, and service levels. "
        "When actual delivery duration differs from the committed duration, teams need a clear way to see the impact before customers, costs, or service commitments are affected. "
        "This project builds a diagnostic analytics layer for delivery performance. It focuses on historical evidence: what happened, where it happened, and which operating dimensions are associated with delay."
    )

    add_heading(document, "2 Problem Statement and Objectives")
    document.add_paragraph(
        "The available order and shipping data contains detailed operational records, but it does not directly provide a consistent view of on-time performance, delay intensity, service-mode efficiency, or geographic risk. "
        "Without this view, logistics teams may respond to individual late shipments instead of identifying repeatable patterns.")
    add_bullets(document, [
        "Measure the proportion of deliveries that are early, on time, or delayed.",
        "Quantify the average difference between actual and scheduled shipping duration.",
        "Compare delivery performance across shipping modes.",
        "Identify high-risk regions, markets, and customer segments.",
        "Deliver findings through an interactive Streamlit dashboard for non-technical stakeholders.",
    ])

    add_heading(document, "3 Dataset and Preparation")
    document.add_paragraph(
        "The supplied APL Logistics dataset contains 180,519 shipment-level records. It includes actual and scheduled shipping days, delivery status, late-delivery-risk flag, shipping mode, market, order region, country, customer segment, sales, profit, and customer-location coordinates. "
        "The preparation pipeline checked duplicates, validated shipping-duration fields, standardized dashboard filter fields, and produced an analytics-ready dataset.")
    add_table(document, ["Data check", "Result"], [
        ["Source rows", "180,519"],
        ["Exact duplicate rows removed", "0"],
        ["Invalid shipping-duration rows removed", "0"],
        ["Analytics-ready rows", "180,519"],
        ["Date fields available", "No order or delivery date field in supplied CSV"],
    ], [3.3, 3.3])

    add_heading(document, "4 Analytical Methodology")
    document.add_paragraph(
        "The core measure is the delivery delay gap. It is calculated as actual shipping days minus scheduled shipping days. A negative result is classified as early, zero as on time, and a positive result as delayed. "
        "The analysis then aggregates the result across operational dimensions. The dashboard also displays the source Late_delivery_risk field separately because it does not exactly match the calculated delivery classification.")
    add_table(document, ["Metric", "Definition"], [
        ["On-time or early rate", "Percentage of shipments with a delay gap less than or equal to zero"],
        ["Calculated delayed rate", "Percentage of shipments with a positive delay gap"],
        ["Average delivery delay gap", "Mean of actual shipping days minus scheduled shipping days"],
        ["Source late-risk ratio", "Percentage of shipments marked as late risk in the supplied data"],
    ], [2.1, 4.5])

    add_heading(document, "5 Findings")
    document.add_paragraph(
        "The overall baseline indicates material delay exposure. The calculated delayed-delivery rate is higher than the source late-risk ratio, confirming that the two measures should not be treated as interchangeable. The dashboard therefore presents both measures with clear labels.")
    add_table(document, ["KPI", "Value"], [
        ["Total shipments", "180,519"],
        ["On-time or early delivery rate", "42.72%"],
        ["Calculated delayed-delivery rate", "57.28%"],
        ["Average delivery delay gap", "0.57 days"],
        ["Source late-delivery-risk ratio", "54.83%"],
    ], [3.3, 3.3])
    document.add_paragraph(
        "Shipping mode is the clearest differentiator. First Class recorded a calculated delayed-delivery rate of 100 percent across 27,814 shipments. Second Class had the largest average delay gap at 1.99 days and a delayed rate of 79.73 percent. Standard Class carried the highest shipment volume and had the strongest on-time-or-early rate among the four modes.")
    add_table(document, ["Shipping mode", "Shipments", "Delayed rate", "Average delay gap", "On-time or early"], [
        ["First Class", "27,814", "100.00%", "1.00 days", "0.00%"],
        ["Second Class", "35,216", "79.73%", "1.99 days", "20.27%"],
        ["Same Day", "9,737", "47.83%", "0.48 days", "52.17%"],
        ["Standard Class", "107,752", "39.77%", "0.00 days", "60.23%"],
    ], [1.35, 1.05, 1.2, 1.3, 1.4])
    document.add_paragraph(
        "Regional results should be evaluated using both rate and volume. Central Africa had the highest regional calculated delay rate at 60.70 percent, but with 1,677 shipments. Western Europe had a lower rate of 58.52 percent but a much larger volume of 27,109 shipments. This makes Western Europe a substantial operational priority even though it is not the highest-rate outlier. Customer-segment delay rates were similar, ranging from 57.11 percent to 57.52 percent, so segment alone does not explain the main delay pattern.")

    add_heading(document, "6 Dashboard Implementation")
    document.add_paragraph(
        "The solution is implemented as a Streamlit dashboard using Python, Pandas, Altair, and PyDeck. Users can filter by shipping mode, market, order region, country, and customer segment. The dashboard includes KPI cards, delivery classification, delay-gap distribution, shipping-mode comparison, regional diagnostics, a geographic delay-risk map, a regional priority matrix, downloadable filtered data, and dynamic recommendations.")

    add_heading(document, "7 Recommendations")
    add_bullets(document, [
        "Review First Class scheduling commitments against actual carrier lead times before presenting the service as a premium expedited option.",
        "Investigate Second Class because its 1.99-day average delay gap is the largest among the shipping modes.",
        "Prioritize high-volume regions such as Western Europe alongside high-rate outliers, using the regional priority matrix to balance impact and risk.",
        "Use the dashboard in operational reviews to monitor filters and investigate combinations of region, shipping mode, and customer segment.",
        "Keep the calculated delay metric and source late-risk flag separate in management reporting.",
    ])

    add_heading(document, "8 Limitations and Future Scope")
    document.add_paragraph(
        "The supplied CSV does not include an order or delivery date, so time-series trends and date-range filtering are outside the current scope. The analysis is diagnostic and identifies historical associations; it does not prove causation or predict future outcomes. A later phase could add time-based data, live operational feeds, predictive delay modeling, and route or shipping-mode optimization after stakeholders validate the diagnostic findings.")

    add_heading(document, "9 Conclusion")
    document.add_paragraph(
        "This project converts a large shipment dataset into an operational decision tool. It establishes a verified baseline for delivery performance, identifies significant shipping-mode risk, and gives stakeholders a direct way to explore geographic and customer-related delay patterns. The current dashboard is a complete diagnostic foundation for improving logistics visibility before any predictive or optimization work is introduced.")
    path = OUTPUT_DIR / "APL_Logistics_Research_Paper.docx"
    document.save(path)
    return path


def write_executive_summary() -> Path:
    document = Document()
    configure_document(document)
    add_title_block(
        document,
        "APL Logistics Delivery Performance Executive Summary",
        "Stakeholder briefing on historical delivery performance",
    )
    document.add_paragraph(
        "The delivery-performance dashboard shows that delay exposure is material and concentrated in specific shipping modes. The immediate priority is to review service promises and operating conditions for First Class and Second Class shipments. The dashboard is ready for managers to filter by market, region, country, and customer segment and to download the underlying filtered records for follow-up.")

    add_heading(document, "What the analysis found")
    add_table(document, ["Headline measure", "Result", "What it means"], [
        ["Total shipment records", "180,519", "A large historical base for operational comparison"],
        ["Calculated delayed deliveries", "57.28%", "More than half of shipments exceeded scheduled duration"],
        ["On-time or early deliveries", "42.72%", "Current baseline for service improvement"],
        ["Average delivery delay gap", "0.57 days", "Actual shipping duration is moderately above plan overall"],
        ["Source late-risk flag", "54.83%", "Useful source indicator but not the same as the calculated delay metric"],
    ], [2.0, 1.0, 3.6])

    add_heading(document, "Where attention is needed")
    document.add_paragraph(
        "First Class is the most urgent issue. All 27,814 First Class shipments were delayed according to the actual-versus-scheduled calculation. Second Class requires separate review because it has the highest average delay gap at 1.99 days. Standard Class carries the largest shipment volume and performs comparatively better, with 60.23 percent of shipments on time or early.")
    document.add_paragraph(
        "Regional decisions should not be made from delay rate alone. Central Africa had the highest rate at 60.70 percent, while Western Europe had 27,109 shipments and a 58.52 percent delayed rate. The regional priority matrix in the dashboard balances volume and delay risk so managers can focus on the largest operational impact.")

    add_heading(document, "Recommended actions")
    add_bullets(document, [
        "Validate whether First Class promised durations are realistic for the carrier and route combinations being used.",
        "Investigate the causes of the Second Class average delay gap before changing service targets.",
        "Use the dashboard before regional operations reviews to isolate high-volume and high-delay combinations.",
        "Export filtered records from the dashboard when a regional or shipping-mode owner needs a detailed follow-up list.",
        "Report calculated delay performance and the source risk flag as separate measures.",
    ])

    add_heading(document, "What stakeholders can use now")
    document.add_paragraph(
        "The Streamlit dashboard is ready for local use. It provides KPI cards, delay distributions, shipping-mode comparison, regional and geographic risk views, a priority matrix, filter controls, and a filtered-data download. It is designed for diagnosis and prioritization, not for predicting future delivery delays.")

    add_heading(document, "Scope note")
    document.add_paragraph(
        "The supplied data has no order or delivery date field, so the dashboard intentionally does not provide date-range filtering or time-trend analysis. A future phase can add those capabilities when time-stamped operational data becomes available.")
    path = OUTPUT_DIR / "APL_Logistics_Executive_Summary.docx"
    document.save(path)
    return path


def main() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)
    paper = write_research_paper()
    summary = write_executive_summary()
    print(paper)
    print(summary)


if __name__ == "__main__":
    main()
